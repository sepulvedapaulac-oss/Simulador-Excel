#!/usr/bin/env python3
"""Genera una tarea de Moodle con rúbrica (evaluación del módulo · Parte 2) desde un solo archivo de datos.

Uso: python3 tools/build_tarea.py src/modulo2/tarea.py [src/modulo3/tarea.py ...]
Salida en dist/moduloN/tarea/:
  Moodle_Evaluacion_Modulo_N_Parte_2.txt               instrucción para la descripción de la Tarea (Código fuente)
  Plantilla_respuesta_Evaluacion_Modulo_N_Parte_2.docx  plantilla que descarga y entrega la persona participante
  Rubrica_Evaluacion_Modulo_N_Parte_2.docx              rúbrica y pauta de corrección para el relator o tutor

El aprendizaje esperado sale de curso.json → "sence" (misma fuente que la planilla SENCE).
Los colores son los del curso (tools/moodle.py → paleta()). Ejemplo de datos: src/ejemplos/tarea/tarea.py.
"""
import runpy, sys
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
from curso import curso  # noqa: E402
from moodle import Moodle, paleta, sence, fmt  # noqa: E402

NIVELES = ["Excelente", "Bueno", "Suficiente", "Insuficiente"]
FACTOR = [1, 0.75, 0.5, 0]


def puntos(peso):
    return [round(peso * f) for f in FACTOR]


def nota(pct, exig):
    """Puntaje (0–100) → nota 1,0–7,0 con la exigencia del curso."""
    return 1 + 3 * pct / exig if pct < exig else 4 + 3 * (pct - exig) / (100 - exig)


def escala(exig):
    pares = sorted({exig, 80, 100})
    return "; ".join(f"{p} puntos = {nota(p, exig):.1f}".replace(".", ",") for p in pares)


# ---------------- Instrucción para Moodle (.txt) ----------------
def instruccion(T, ae, exig, rel):
    n, C = T["MODULO"], curso()
    m = Moodle()
    m.hero(f"Evaluación del Módulo {n} · Parte 2", f"«{T['TITULO']}»", f"{C['nombre']} · Evaluación con rúbrica, la califica tu {rel}")
    for x in T["INTRO"]:
        m.p(x)
    m.p(f"Esta es la **Parte 2 de la evaluación del Módulo {n}** (la Parte 1 es el cuestionario). La revisa tu {rel} con una rúbrica y te devuelve retroalimentación escrita.")
    R = T["RUBRICA"]
    m.ficha("Antes de empezar", [
        ("Aprendizaje esperado", ae),
        ("Duración", f"{T['DURACION']} minutos. Puedes hacerla en más de una sesión: tus avances quedan guardados en la plantilla."),
        ("Plazo", "La fecha límite que aparece en esta actividad."),
        ("Qué necesitas", "El caso (más abajo) y la **plantilla de respuesta**, que descargas al final de estas instrucciones."),
        ("Qué entregas", "La plantilla completa, en Word o PDF. Si prefieres, puedes escribir tus respuestas en el cuadro de texto de la entrega."),
        ("Calificación", f"Rúbrica de {len(R)} criterios (100 puntos). Con {exig} puntos obtienes nota 4,0."),
    ])
    m.h4("Paso a paso")
    m.pasos(list(T["PASOS"]) + [("Revisa y entrega.", f"Compara tu trabajo con los criterios de evaluación (más abajo). Guarda el archivo con tu nombre, "
                                                     f"por ejemplo «Evaluacion_M{n}_Nombre_Apellido», y súbelo con el botón «Agregar entrega». Luego presiona «Guardar cambios».")])
    if T.get("HERRAMIENTAS"):
        m.h4("Herramientas del módulo que vas a usar")
        m.lista(T["HERRAMIENTAS"])
    m.h4("Criterios de evaluación")
    m.tabla(["Criterio", "Peso", "Para lograr el nivel más alto"], [(f"{i}. {r['criterio']}", f"{fmt(r['peso'])} %", r["alto"]) for i, r in enumerate(R, 1)])
    m.p("La rúbrica completa, con los cuatro niveles de cada criterio, está al final de estas instrucciones.")
    K = T["CASO"]
    m.h3(f"El caso: «{K['titulo']}»")
    for b in K["bloques"]:
        m.tabla(b["cols"], b["filas"]) if isinstance(b, dict) else m.p(b)
    m.h3("Rúbrica de evaluación")
    m.p("Cada criterio se evalúa en cuatro niveles. El puntaje de cada nivel ya considera la ponderación del criterio; la suma máxima es 100 puntos.")
    for i, r in enumerate(R, 1):
        m.rubrica(i, r["criterio"], r["peso"], list(zip(NIVELES, puntos(r["peso"]), r["niveles"])))
    m.p(f"**Nota de la evaluación:** el puntaje (0 a 100) se convierte a nota de 1,0 a 7,0 con {exig} % de exigencia: {escala(exig)}.")
    m.cierre("Descarga la plantilla de respuesta aquí abajo ⬇",
             f"Complétala, guárdala con tu nombre y súbela con «Agregar entrega». Si tienes dudas, escribe a tu {rel} por la mensajería del aula virtual.")
    return m


# ---------------- Word ----------------
def hx(c):
    return c.lstrip("#").upper()


def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hx(color))
    tcpr.append(shd)


def texto(par, t, size=10, color=None, bold=False, italic=False):
    for i, part in enumerate(t.split("**")):
        if not part:
            continue
        r = par.add_run(part)
        r.bold = bold or i % 2 == 1
        r.italic = italic
        r.font.size = Pt(size); r.font.name = "Arial"
        if color:
            r.font.color.rgb = RGBColor.from_string(hx(color))
    return par


def doc_base():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.8)
    s.top_margin = s.bottom_margin = Cm(1.8)
    st = d.styles["Normal"]
    st.font.name = "Arial"; st.font.size = Pt(10)
    return d


def p(d, t, **kw):
    par = d.add_paragraph()
    par.paragraph_format.space_after = Pt(kw.pop("after", 6))
    return texto(par, t, **kw)


def tabla(d, cols, filas, P, widths=None, ejemplo=None, alto=0.9):
    """Tabla con encabezado oscuro, fila de ejemplo sombreada (✎) y filas en blanco."""
    t = d.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(cols):
        c = t.rows[0].cells[i]
        shade(c, P["oscuro"])
        texto(c.paragraphs[0], h, size=9.5, color="#FFFFFF", bold=True)
    if ejemplo:
        r = t.add_row().cells
        for i, v in enumerate(ejemplo):
            shade(r[i], P["suave"])
            texto(r[i].paragraphs[0], ("✎ " if i == 0 else "") + v, size=9, color=P["gris"], italic=True)
    for f in filas:
        r = t.add_row()
        r.height = Cm(alto)
        for i, v in enumerate(f if isinstance(f, (list, tuple)) else [f] + [""] * (len(cols) - 1)):
            texto(r.cells[i].paragraphs[0], v, size=9.5)
    if widths:
        tot = sum(widths)
        t.autofit = False
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(17.4 * w / tot)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(17.4 * w / tot)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def titulo(d, t, sub, P):
    texto(d.add_paragraph(), t, size=16, color=P["oscuro"], bold=True)
    p(d, sub, size=10.5, color=P["primario"], after=10)


def plantilla(T, ae, P):
    n, C = T["MODULO"], curso()
    d = doc_base()
    titulo(d, f"Plantilla de respuesta · Evaluación del Módulo {n} · Parte 2", f"{T['TITULO']} · {C['nombre']}", P)
    p(d, "**Cómo usar esta plantilla:** escribe tus respuestas en las filas en blanco de cada tabla. La fila sombreada (✎) es un ejemplo resuelto: "
         "muestra el nivel de detalle esperado y no cuenta como respuesta tuya. Las celdas crecen a medida que escribes. Guarda el archivo con tu nombre y súbelo a la tarea.")
    tabla(d, ["Dato", "Completa aquí"], [["Nombre", ""], ["Fecha de entrega", ""]], P, widths=[30, 70], alto=0.7)
    for s in T["PLANTILLA"]:
        texto(d.add_paragraph(), s["titulo"], size=12, color=P["primario"], bold=True)
        p(d, s["inst"], size=9.5)
        if s.get("cols"):
            tabla(d, s["cols"], [[""] * len(s["cols"]) for _ in range(s.get("filas", 3))], P, widths=s.get("anchos"), ejemplo=s.get("ejemplo"), alto=s.get("alto", 1.4))
        else:
            tabla(d, [s.get("encabezado", "Respuesta")], list(s.get("lineas", [])) + [""], P, alto=s.get("alto", 6))
    p(d, f"Aprendizaje esperado que evidencias: {ae}", size=8.5, color=P["gris"], italic=True)
    return d


def rubrica_doc(T, ae, crit, exig, rel, P):
    n, C = T["MODULO"], curso()
    d = doc_base()
    titulo(d, f"Rúbrica y pauta de corrección · Evaluación del Módulo {n} · Parte 2", f"{T['TITULO']} · {C['nombre']} · Documento para el {rel}", P)
    tabla(d, ["Dato", "Detalle"], [["Aprendizaje esperado", ae], ["Criterio de evaluación (planilla SENCE)", crit],
                                   ["Puntaje y nota", f"100 puntos. Nota 4,0 con {exig} puntos ({exig} % de exigencia). {escala(exig)}."],
                                   ["Retroalimentación", T.get("PLAZO", "Escrita, en un plazo máximo de 48 horas hábiles desde la entrega: qué logró, qué corregir y cómo mejorar.")]],
          P, widths=[32, 68], alto=0.6)
    texto(d.add_paragraph(), "Rúbrica", size=12, color=P["primario"], bold=True)
    filas = [[f"{i}. {r['criterio']} ({fmt(r['peso'])} %)"] + [f"**{fmt(pt)} pts.** {x}" for pt, x in zip(puntos(r["peso"]), r["niveles"])]
             for i, r in enumerate(T["RUBRICA"], 1)]
    tabla(d, ["Criterio"] + NIVELES, filas, P, widths=[22, 20, 20, 19, 19], alto=0.6)
    if T.get("PAUTA"):
        texto(d.add_paragraph(), "Pauta: respuestas esperadas", size=12, color=P["primario"], bold=True)
        for x in T["PAUTA"]:
            par = d.add_paragraph(style="List Bullet")
            texto(par, x, size=9.5)
    texto(d.add_paragraph(), "Registro de corrección", size=12, color=P["primario"], bold=True)
    tabla(d, ["Criterio", "Nivel", "Puntos", "Comentario"], [[f"{i}. {r['criterio']}", "", "", ""] for i, r in enumerate(T["RUBRICA"], 1)] + [["Total y nota", "", "", ""]],
          P, widths=[34, 14, 10, 42], alto=1.0)
    return d


def build(src):
    src = (ROOT / src).resolve()
    T = runpy.run_path(str(src))
    n = T["MODULO"]
    S = sence()
    ae, crit = S["ae"][n - 1], S["criterios"][n - 1]
    C = curso()
    exig = C.get("evaluacion", {}).get("exigencia", 60)
    rel = C.get("relator", "relator o tutor")
    assert abs(sum(r["peso"] for r in T["RUBRICA"]) - 100) < 1e-6, "Los pesos de la rúbrica deben sumar 100"
    assert all(len(r["niveles"]) == 4 for r in T["RUBRICA"]), "Cada criterio necesita 4 niveles"
    P = paleta()
    out = ROOT / "dist" / f"modulo{n}" / "tarea"
    out.mkdir(parents=True, exist_ok=True)
    instruccion(T, ae, exig, rel).guardar(out / f"Moodle_Evaluacion_Modulo_{n}_Parte_2.txt")
    for nombre, doc in ((f"Plantilla_respuesta_Evaluacion_Modulo_{n}_Parte_2.docx", plantilla(T, ae, P)),
                        (f"Rubrica_Evaluacion_Modulo_{n}_Parte_2.docx", rubrica_doc(T, ae, crit, exig, rel, P))):
        doc.save(out / nombre)
        print(f"OK {(out / nombre).relative_to(ROOT)}")


if __name__ == "__main__":
    for f in sys.argv[1:]:
        build(f)
